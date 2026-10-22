# -*- coding: utf-8 -*-
"""生成《课程知识与技能清单》Word 文档"""
from docx import Document
from docx.shared import Pt, RGBColor
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document()

# 全局中文字体
style = doc.styles['Normal']
style.font.name = 'Calibri'
style.font.size = Pt(11)
style.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')

def set_cn(run, name='微软雅黑'):
    run.font.name = name
    run._element.rPr.rFonts.set(qn('w:eastAsia'), name)

def title(text):
    p = doc.add_heading(text, level=0)
    for r in p.runs:
        set_cn(r, '微软雅黑')
        r.font.size = Pt(22)
        r.font.color.rgb = RGBColor(0, 0, 0)   # 黑色
    return p

def h1(text):
    p = doc.add_heading(text, level=1)
    for r in p.runs:
        set_cn(r, '微软雅黑')
        r.font.color.rgb = RGBColor(0, 0, 0)   # 黑色
    return p

def h2(text):
    p = doc.add_heading(text, level=2)
    for r in p.runs:
        set_cn(r, '微软雅黑')
        r.font.color.rgb = RGBColor(0, 0, 0)   # 黑色
    return p

def bullet(text, bold_prefix=None):
    p = doc.add_paragraph(style='List Bullet')
    if bold_prefix:
        r1 = p.add_run(bold_prefix)
        r1.bold = True
        set_cn(r1)
    r = p.add_run(text)
    set_cn(r)
    return p

def para(text, bold=False):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = bold
    set_cn(r)
    return p

# ============ 封面 ============
title('2024 电赛 H 题小车 —— 课程知识与技能清单')
p = doc.add_paragraph()
r = p.add_run('(基于 STM32F103 自动行驶小车项目的完整开发过程整理)')
r.italic = True
set_cn(r)
p.alignment = WD_ALIGN_PARAGRAPH.CENTER

para('本清单覆盖:STM32 嵌入式基础、传感器与外设、控制算法、软件工程、系统调试、'
     'AI 辅助开发流程、工程安全经验与比赛全流程。所有条目均来自本项目实际做过、'
     '调过、踩过的内容。')

# ============ 一、STM32 基础 ============
h1('一、STM32 嵌入式基础知识')

h2('1.1 芯片与开发环境')
bullet('STM32F103(Cortex-M3)最小系统结构:供电、时钟、复位、BOOT 配置', '芯片结构:')
bullet('CubeMX 初始化代码的结构与 USER CODE 区域的作用(重新生成代码不丢自定义内容)')
bullet('HAL 库的基本使用模式:句柄(HandleTypeDef)、初始化(Init)、回调(Callback)')
bullet('GCC + CMake + Ninja 命令行编译流程,以及编译输出解读(警告/错误/链接成功)')
bullet('烧录与调试工具(ST-Link SWD)的正确使用方法')

h2('1.2 GPIO 与外设')
bullet('GPIO 的输入/输出模式、推挽/开漏、上拉/下拉的选用场景(如 I2C 用开漏)')
bullet('按键输入:外部中断(EXTI)检测、软件消抖、长按/短按区分、按键状态机')
bullet('LED、蜂鸣器等开关量控制;三极管驱动的电平逻辑(高有效/低有效)')
bullet('引脚复用与调试口冲突:JTAG/SWD 重映射(REMAP_SWJ_NOJTAG)释放 PA15 等引脚')

h2('1.3 定时器')
bullet('PWM 输出:频率与占空比的设定,驱动电机调速的原理')
bullet('定时器编码器接口模式:正交解码、四倍频计数,用于轮速测量')
bullet('周期定时中断(如 10ms 控制周期)作为所有控制逻辑的节拍')
bullet('AT8236 电机驱动桥:前进/后退/刹车/滑行四种状态的实现方式')

h2('1.4 串口与通信')
bullet('USART 基本配置:波特率、数据位、校验、流控')
bullet('三种接收方式对比:轮询(简单但会溢出)、阻塞(可靠但占 CPU)、中断/DMA(推荐)')
bullet('printf 重定向到串口(__io_putchar)用于调试输出')
bullet('串口接收溢出(ORE)的成因与排查:发送快于读取时丢字节,教训:中断接收才是正道')
bullet('I2C 协议时序:软件模拟 I2C 的起始/停止/应答/读写的实现')

h2('1.5 中断系统')
bullet('NVIC 中断优先级的配置与抢占概念')
bullet('外部中断(EXTI)与串口中断服务函数的写法(IRQHandler → HAL 回调)')
bullet('中断与主循环的数据交换:volatile 标志位的典型用法')

# ============ 二、传感器与外设 ============
h1('二、传感器与外设驱动')

h2('2.1 四路循迹传感器')
bullet('数字量传感器的读取与滤波;黑色/白色电平约定(0/1)与传感器物理排列的对应关系')
bullet('窄黑线工况的典型现象:同一时刻只有 1~2 路压线、传感器之间存在"死区缝隙"')

h2('2.2 编码器')
bullet('编码器计数与转速(RPM)换算:计数/圈 × 时间的关系')
bullet('编码器计距:转速积分得到里程,用于任务的距离窗口与标定')
bullet('轮周长标定方法:推车实测距离反推每圈计数对应的毫米数')

h2('2.3 MPU6050 陀螺仪')
bullet('MPU6050 的原始数据读取与单位换算(LSB 与 °/s)')
bullet('两种姿态方案对比:DMP 库(黑盒,依赖其内部校准)与原始数据自己积分(可控)')
bullet('陀螺仪零偏校准:上电静止采样取平均;预热期的必要性(零偏随温度漂移)')
bullet('航向角(yaw)积分与死区处理;无磁力计时 yaw 只能相对、会漂移的物理局限')
bullet('符号约定实验:手动转车看读数,确定"正角度=左转还是右转"')

h2('2.4 OLED 显示')
bullet('软件 SPI 驱动 OLED 的初始化序列与显示原理(页/列寻址)')
bullet('自建 5×7 点阵字库、字符查表、显示函数封装')
bullet('用 OLED 做实时调试面板:传感器值、转速、航向、状态机阶段、故障码')

# ============ 三、控制算法 ============
h1('三、控制理论与算法')

h2('3.1 PID 控制')
bullet('P/I/D 三项的物理含义与各自作用:比例响应、积分消除稳态偏差、微分阻尼')
bullet('增量式与位置式 PID 的区别与适用场景')
bullet('循迹 PID:传感器加权误差、误差归一化、积分限幅、输出限幅')
bullet('航向保持 PID:用航向角偏差做差速修正,参数温和(小增益、小限幅)')
bullet('右轮追左轮的速度同步 PID:误差=左轮转速−右轮转速,配轮径补偿')
bullet('PID 参数整定方法:先 P 后 I 再 D、单参数单次改动、现象对照调参表')

h2('3.2 有限状态机(FSM)')
bullet('任务级状态机:直行→循迹→直行→循迹 的段间切换与触发条件(路口/线尾)')
bullet('按键循环状态机:待机/运行/停止的循环切换,自然完成与手动急停')
bullet('阶段状态机(任务3):转向→静止确认→斜线→顶点制动→捕捉→弧线→终点')
bullet('状态机设计要点:每个状态职责单一、切换条件明确、进入状态时初始化')

h2('3.3 滤波与信号处理')
bullet('滑动平均滤波、一阶低通(IIR)滤波的写法与响应速度权衡')
bullet('死区(deadband)的作用:消除量化噪声导致的控制抖动')
bullet('去抖/确认机制:信号持续 N 毫秒才采信,防止毛刺误触发')

h2('3.4 容错与恢复策略')
bullet('恢复模式:线跑到最外侧时的"急刹→减速追线→压稳退出"流程')
bullet('丢线间隙保持:短暂全白时保持上次差速,不误判为丢线')
bullet('超距救援:规定距离没压线时不判故障,摆正后直行找线,设最终兜底阈值')
bullet('超时保护:转向/捕捉/总时间的多级超时与故障码(便于定位)')
bullet('多级兜底设计思想:主方案 + 备用方案(陀螺仪失效退回编码器同步)')

# ============ 四、软件工程 ============
h1('四、软件工程方法')

h2('4.1 代码组织')
bullet('模块化:每个功能一个 .c/.h(motor/track/oled/mpu6050/h_course)')
bullet('参数集中管理:所有可调参数放头文件并配中文注释,调参不改逻辑')
bullet('命名规范:函数名/宏名/变量名的统一风格与含义自解释')
bullet('注释习惯:每个参数写"管什么、什么时候调、调大调小会怎样"')

h2('4.2 调试系统')
bullet('多级调试手段:OLED 实时面板、蜂鸣器/LED 事件提示、串口调试开关')
bullet('故障码设计:每个故障给编号,现场一报码就定位')
bullet('对照实验法:关掉某个环节(如航向保持)跑一次,二分定位问题')

h2('4.3 版本管理')
bullet('git 的基本使用:提交、查看历史、回滚')

# ============ 五、AI 辅助开发 ============
h1('五、AI 辅助开发流程(课程重点)')

h2('5.1 AI 生成代码的组成与工作方式')
bullet('需求→生成:把目标行为描述清楚(如"任意一路压黑就停"),AI 产出对应代码')
bullet('代码审查:AI 生成的代码必须逐段读懂再上板,重点检查符号、单位、边界条件')
bullet('参考代码学习:把成熟方案(如 Sunward)交给 AI 分析,提取可复用的架构')

h2('5.2 与 AI 的协作迭代方法')
bullet('现象描述法:反馈必须带数据(屏幕读数、故障码、偏转方向),而非"不好使"')
bullet('参数整定闭环:现象→AI 分析→改参数→复测,每轮只动一个变量')
bullet('诊断性实验设计:如"手动转车看 Y 值"一次定死符号约定')
bullet('让 AI 写调试工具:距离标定模式、故障码显示等辅助手段')

h2('5.3 AI 协作中的关键教训')
bullet('符号约定与物理量纲必须实测确认(本项目多次踩坑:yaw 方向、电机左右映射)')
bullet('AI 的修改要可回退:每次改动小而独立,错了能一键还原')
bullet('AI 无法替代硬件排查:ST-Link 烧毁是电源问题,不是代码问题——学会分辨问题层级')
bullet('AI 的建议需要结合场地实测:几何公式算的和场地胶带贴的往往有差异')

h2('5.4 提示词与文档能力')
bullet('需求翻译:把"感觉跑偏"翻译成可执行规格(方向/大小/条件)')
bullet('让 AI 产出工程文档:汇报稿、知识清单、调参手册')

# ============ 六、工程与安全 ============
h1('六、工程实践与安全意识')

h2('6.1 供电与接线安全')
bullet('双电源顶牛的危害:电池与调试器 3.3V 同时供电导致过热烧毁的教训')
bullet('正确供电习惯:烧录断电池、跑车拔调试器,或永久断开调试器 3.3V 线')
bullet('万用表的基本使用:测短路、测电压、测电阻判断硬件好坏')

h2('6.2 标定与测试流程')
bullet('场地实测:几何角度、线头位置、轮周长都要实测而非照搬图纸')
bullet('分阶段验证:一段一段跑,用状态显示判断每段是否通过')
bullet('可重复性:同样参数多次跑,区分"偶发"与"必现"问题')

# ============ 七、比赛流程 ============
h1('七、比赛与汇报能力')

bullet('赛前准备:备件清单、工具清单、程序版本确认')
bullet('场地适应:到新场地先标定(轮周长、角度、线宽)再参赛')
bullet('技术汇报:能把控制架构讲清楚(传感器→控制→执行的三层结构)')
bullet('答辩问答:原理性问题(为什么不用 DMP、PID 各参数作用)的准备')

# ============ 八、通用能力 ============
h1('八、可迁移的通用能力')

bullet('问题分解:把"车跑不好"拆成 方向/速度/判定/硬件 四层逐层排查')
bullet('数据驱动的调试:先看数据(屏幕/串口),再动手改')
bullet('控制直觉:理解惯性、过冲、稳态偏差、正反馈等概念')
bullet('工程文档习惯:参数注释、调参记录、故障记录')
bullet('协作沟通:与 AI、与队友的准确描述问题的能力')

doc.save(r'C:\Users\hshsgxlj\Desktop\2024电赛小车\stm32\stm32\课程知识与技能清单.docx')
print('OK saved')
